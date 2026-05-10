#!/usr/bin/env python3
"""
Homeschool Budget Spreadsheet Generator
Generates customizable Excel/Google Sheets budget templates for homeschooling families
Supports monthly, annual, and category-specific budgets
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

class BudgetSpreadsheetGenerator:
    def __init__(self, output_dir="./budget_templates"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

        # Styles
        self.header_font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
        self.header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        self.total_font = Font(name='Arial', size=12, bold=True)
        self.total_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    def create_monthly_budget(self, family_name="Homeschool Family", year=2025):
        """Create a comprehensive monthly homeschool budget"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Monthly Budget"

        # Title
        ws.merge_cells('A1:F1')
        ws['A1'] = f"Homeschool Monthly Budget - {family_name}"
        ws['A1'].font = Font(name='Arial', size=16, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.merge_cells('A2:F2')
        ws['A2'] = f"Year: {year}"
        ws['A2'].font = Font(italic=True)
        ws['A2'].alignment = Alignment(horizontal='center')

        # Headers
        headers = ['Category', 'Item Description', 'Budgeted', 'Actual', 'Difference', 'Notes']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center')

        # Categories and sample items
        categories = [
            ('Curriculum', [
                'Math curriculum',
                'Reading/Literature',
                'Science materials',
                'History/Geography',
                'Art supplies',
                'Music lessons'
            ]),
            ('Extracurricular', [
                'Sports fees',
                'Club dues',
                'Field trips',
                'Swim/art classes'
            ]),
            ('Supplies', [
                'Paper, pencils, notebooks',
                'Printer ink',
                'Craft supplies',
                'Educational toys'
            ]),
            ('Technology', [
                'Educational software',
                'Tablet/laptop',
                'Internet'
            ]),
            ('Testing/Assessment', [
                'Standardized tests',
                'Portfolio review'
            ]),
            ('Miscellaneous', [
                'Co-op fees',
                'Library fines',
                'Other'
            ])
        ]

        row = 5
        total_budget = 0
        total_actual = 0

        for category, items in categories:
            # Category header
            ws.cell(row=row, column=1, value=category).font = Font(bold=True)
            ws.cell(row=row, column=2, value='').font = Font(bold=True)
            ws.row_dimensions[row].height = 20
            row += 1

            for item in items:
                ws.cell(row=row, column=1, value='')  # empty category col
                ws.cell(row=row, column=2, value=item)
                ws.cell(row=row, column=3, value=0.0)  # Budgeted
                ws.cell(row=row, column=4, value=0.0)  # Actual
                ws.cell(row=row, column=5, value='=D{row}-C{row}'.format(row=row))
                ws.cell(row=row, column=6, value='')
                row += 1

        # Totals section
        total_row = row + 2
        ws.cell(row=total_row, column=2, value='TOTAL BUDGET').font = self.total_font
        ws.cell(row=total_row, column=3, value=f'=SUM(C5:C{row-1})').font = self.total_font
        ws.cell(row=total_row, column=3).fill = self.total_fill

        ws.cell(row=total_row+1, column=2, value='TOTAL ACTUAL').font = self.total_font
        ws.cell(row=total_row+1, column=4, value=f'=SUM(D5:D{row-1})').font = self.total_font
        ws.cell(row=total_row+1, column=4).fill = self.total_fill

        ws.cell(row=total_row+2, column=2, value='NET DIFFERENCE').font = Font(bold=True)
        ws.cell(row=total_row+2, column=5, value=f'=D{total_row+1}-C{total_row}').font = Font(bold=True)

        # Formatting: currency format for numeric columns
        for col in [3, 4, 5]:  # Budgeted, Actual, Difference
            col_letter = get_column_letter(col)
            for r in range(5, row):
                ws[f'{col_letter}{r}'].number_format = '"$"#,##0.00'

        # Column widths
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 20

        filename = f"{self.output_dir}/homeschool_monthly_budget_{family_name.replace(' ', '_')}.xlsx"
        wb.save(filename)
        return filename

    def create_annual_budget(self, family_name="Homeschool Family", year=2025):
        """Create an annual budget overview with monthly breakdown"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Annual Budget"

        # Title
        ws.merge_cells('A1:M1')
        ws['A1'] = f"Homeschool Annual Budget - {family_name}"
        ws['A1'].font = Font(name='Arial', size=16, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.merge_cells('A2:M2')
        ws['A2'] = f"Year: {year}"
        ws['A2'].font = Font(italic=True)
        ws['A2'].alignment = Alignment(horizontal='center')

        # Headers
        ws.cell(row=4, column=1, value='Category').font = self.header_font
        ws.cell(row=4, column=1).fill = self.header_fill

        months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun',
                  'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
        for col, month in enumerate(months, 2):
            ws.cell(row=4, column=col, value=month).font = self.header_font
            ws.cell(row=4, column=col).fill = self.header_fill

        ws.cell(row=4, column=14, value='Annual Total').font = self.header_font
        ws.cell(row=4, column=14).fill = self.header_fill

        # Categories (matching monthly template)
        categories = [
            'Curriculum',
            'Extracurricular',
            'Supplies',
            'Technology',
            'Testing/Assessment',
            'Miscellaneous',
            'TOTAL'
        ]

        for row, category in enumerate(categories, 5):
            ws.cell(row=row, column=1, value=category)
            if category == 'TOTAL':
                ws.cell(row=row, column=1).font = self.total_font

            # Monthly columns (B-M)
            for col in range(2, 14):
                if category == 'TOTAL':
                    # Sum the column above
                    col_letter = get_column_letter(col)
                    ws.cell(row=row, column=col, value=f'=SUM({col_letter}5:{col_letter}{row-1})')
                    ws.cell(row=row, column=col).font = self.total_font
                else:
                    ws.cell(row=row, column=col, value=0.0)

            # Annual total column (N)
            if category == 'TOTAL':
                ws.cell(row=row, column=14, value=f'=SUM(B{row}:M{row})')
                ws.cell(row=row, column=14).font = self.total_font
            else:
                ws.cell(row=row, column=14, value=f'=SUM(B{row}:M{row})')

        # Formatting: currency for all numeric cells
        for row in range(5, 13):
            for col in range(2, 15):
                ws.cell(row=row, column=col).number_format = '"$"#,##0.00'

        # Column widths
        ws.column_dimensions['A'].width = 20
        for col in range(2, 15):
            ws.column_dimensions[get_column_letter(col)].width = 12

        filename = f"{self.output_dir}/homeschool_annual_budget_{family_name.replace(' ', '_')}.xlsx"
        wb.save(filename)
        return filename

    def create_curriculum_tracker(self, family_name="Homeschool Family", subjects=None):
        """Create a curriculum expense tracker"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Curriculum Tracker"

        if subjects is None:
            subjects = [
                ('Math', 'Singapore Math, Art of Problem Solving'),
                ('Reading', 'Bob Books, Reading Eggs'),
                ('Science', 'Apologia, Mystery Science'),
                ('History', 'Story of the World, Kahn Academy'),
                ('Art', 'Art for Kids Hub, Art lessons')
            ]

        # Title
        ws.merge_cells('A1:E1')
        ws['A1'] = f"Curriculum Purchase Tracker - {family_name}"
        ws['A1'].font = Font(name='Arial', size=16, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center')

        # Headers
        headers = ['Subject', 'Curriculum Name', 'Vendor', 'Price', 'Grade Level', 'Status', 'Notes']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill

        # Sample rows
        row = 4
        for subject, examples in subjects:
            ws.cell(row=row, column=1, value=subject)
            ws.cell(row=row, column=2, value=examples.split(',')[0].strip())
            ws.cell(row=row, column=3, value='')
            ws.cell(row=row, column=4, value=0.0)
            ws.cell(row=row, column=4).number_format = '"$"#,##0.00'
            ws.cell(row=row, column=5, value='')
            ws.cell(row=row, column=6, value='Planned')
            ws.cell(row=row, column=7, value='')
            row += 1

        # Column widths
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 25

        filename = f"{self.output_dir}/curriculum_tracker_{family_name.replace(' ', '_')}.xlsx"
        wb.save(filename)
        return filename

    def generate_full_homeschool_bundle(self, family_name="Homeschool Family", year=2025):
        """Generate a complete bundle: monthly budget, annual overview, curriculum tracker"""
        files = []
        files.append(self.create_monthly_budget(family_name, year))
        files.append(self.create_annual_budget(family_name, year))
        files.append(self.create_curriculum_tracker(family_name))

        return files

if __name__ == "__main__":
    generator = BudgetSpreadsheetGenerator()

    # Test generate a bundle
    bundle_files = generator.generate_full_homeschool_bundle("Test Family")

    print(f"Generated {len(bundle_files)} budget spreadsheets:")
    for f in bundle_files:
        print(f"  - {f}")

    print(f"\nAll files saved to: {generator.output_dir}")